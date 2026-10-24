"""Text extraction from uploaded resumes. Every document is generated in the test."""

import io

import fitz
import pytest
from docx import Document

from app.services.document_parser import DocumentParserService


@pytest.fixture
def parser():
    return DocumentParserService()


async def test_text_and_markdown_bytes_are_decoded(parser):
    assert (
        await parser.extract_text("notes.md", "# Asha Example\nPython".encode())
        == "# Asha Example\nPython"
    )


async def test_text_file_is_read_from_disk(parser, tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("Ravi Sample, backend engineer", encoding="utf-8")
    assert await parser.extract_text(str(path)) == "Ravi Sample, backend engineer"


async def test_docx_paragraphs_and_tables_are_extracted(parser):
    doc = Document()
    doc.add_paragraph("Asha Example")
    doc.add_paragraph("Software Engineer Intern at Sample Labs")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "FastAPI"
    buffer = io.BytesIO()
    doc.save(buffer)

    text = await parser.extract_text("resume.docx", buffer.getvalue())
    assert "Software Engineer Intern at Sample Labs" in text
    assert "Python" in text and "FastAPI" in text


async def test_pdf_text_is_extracted_from_every_page(parser):
    pdf = fitz.open()
    for line in ("Asha Example", "Projects: job-scout"):
        page = pdf.new_page()
        page.insert_text((72, 72), line)
    data = pdf.tobytes()
    pdf.close()

    text = await parser.extract_text("resume.PDF", data)
    assert "Asha Example" in text
    assert "Projects: job-scout" in text


async def test_unsupported_extension_is_rejected(parser):
    with pytest.raises(ValueError, match="Unsupported file type: .exe"):
        await parser.extract_text("resume.exe", b"MZ")


def test_file_hash_is_sha256(parser):
    assert (
        parser.compute_file_hash(b"abc")
        == "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"
    )
